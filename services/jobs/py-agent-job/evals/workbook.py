"""Planilhas editáveis de casos e resultados da última execução."""

import csv
import os
from datetime import UTC, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CASES_DIR = Path(__file__).resolve().parent / "cases"
CASE_HEADERS = ("id", "question", "reference", "mode")
METRICS = ("correcao_factual", "completude", "relevancia", "aderencia_instrucoes", "clareza")
RESULT_HEADERS = (
    "id",
    "executado_em_utc",
    "status",
    "pergunta",
    "referencia",
    "resposta",
    "modo",
    *METRICS,
    "nota_media",
    "tempo_segundos",
    "motivo",
    "sugestao_correcao",
    "erro",
)


def workbook_path(agent_id: str) -> Path:
    return CASES_DIR / f"{agent_id}.xlsx"


def _style_header(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32


def _save_atomic(book: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(suffix=".xlsx", dir=path.parent, delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        book.save(temp_path)
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)


def _refresh_csv_if_present(book: Workbook, agent_id: str) -> None:
    """Mantém um CSV já solicitado em sincronia com a última rodada."""
    target = workbook_path(agent_id).with_suffix(".csv")
    if not target.exists():
        return
    cases = book["Casos"]
    metrics = book["Métricas"]
    results = {}
    for row in metrics.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        values = list(row[1:])
        scores = [score for score in values[6:11] if isinstance(score, (int, float))]
        values[11] = round(sum(scores) / len(scores), 2) if scores else ""
        results[row[0]] = values
    with NamedTemporaryFile(
        mode="w", encoding="utf-8-sig", newline="", suffix=".csv", dir=target.parent, delete=False
    ) as temp:
        temp_path = Path(temp.name)
        writer = csv.writer(temp)
        writer.writerow([cell.value for cell in cases[1]] + [cell.value for cell in metrics[1]][1:])
        for row in cases.iter_rows(min_row=2, max_col=4, values_only=True):
            if row[0]:
                writer.writerow(list(row) + list(results.get(row[0], [""] * (len(RESULT_HEADERS) - 1))))
    try:
        os.replace(temp_path, target)
    finally:
        temp_path.unlink(missing_ok=True)


def create_workbook(agent_id: str, cases: list[dict]) -> Path:
    """Cria a planilha inicial; nunca substitui uma planilha editada pelo usuário."""
    path = workbook_path(agent_id)
    if path.exists():
        return path
    book = Workbook()
    sheet = book.active
    sheet.title = "Casos"
    sheet.append(CASE_HEADERS)
    for case in cases:
        sheet.append(tuple(case.get(key, "") for key in CASE_HEADERS))
    for column, width in {"A": 30, "B": 55, "C": 70, "D": 16}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _style_header(sheet)
    _reset_output_sheets(book)
    _save_atomic(book, path)
    return path


def read_cases(agent_id: str) -> list[dict]:
    path = workbook_path(agent_id)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Casos"]
        if tuple(cell.value for cell in sheet[1][:4]) != CASE_HEADERS:
            raise ValueError(f"Cabeçalhos inválidos na aba Casos de {path}")
        cases = []
        ids = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=4, values_only=True), 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            case = dict(zip(CASE_HEADERS, row))
            if any(not isinstance(case[key], str) or not case[key].strip() for key in CASE_HEADERS):
                raise ValueError(f"Preencha id, question, reference e mode em Casos!{row_number}")
            if case["mode"] not in {"exact", "contains", "semantic"}:
                raise ValueError(f"Modo inválido em Casos!D{row_number}: {case['mode']}")
            if case["id"] in ids:
                raise ValueError(f"ID duplicado em Casos!A{row_number}: {case['id']}")
            ids.add(case["id"])
            cases.append(case)
        return cases
    finally:
        book.close()


def _reset_output_sheets(book: Workbook) -> None:
    for name in ("Métricas", "Resumo"):
        if name in book:
            del book[name]
    sheet = book.create_sheet("Métricas")
    sheet.append(RESULT_HEADERS)
    for index, width in enumerate((28, 23, 14, 50, 60, 70, 15, 18, 16, 16, 22, 12, 14, 18, 55, 55, 55), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    _style_header(sheet)

    summary = book.create_sheet("Resumo")
    summary.append(("Métrica", "Valor"))
    summary.append(("Casos executados", "=COUNTA('Métricas'!A2:A1048576)"))
    summary.append(("Aprovados", "=COUNTIF('Métricas'!C2:C1048576,\"APROVADO\")"))
    summary.append(("Reprovados", "=COUNTIF('Métricas'!C2:C1048576,\"REPROVADO\")"))
    summary.append(("Erros", "=COUNTIF('Métricas'!C2:C1048576,\"ERRO\")"))
    summary.append(("Taxa de aprovação", "=IFERROR(B3/B2,0)"))
    summary["B6"].number_format = "0.0%"
    for offset, metric in enumerate(METRICS, 8):
        letter = get_column_letter(offset)
        summary.append((f"Média {metric}", f"=IFERROR(AVERAGE('Métricas'!{letter}2:{letter}1048576),0)"))
    summary.append(("Tempo médio (s)", "=IFERROR(AVERAGE('Métricas'!N2:N1048576),0)"))
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 22
    _style_header(summary)


def reset_results(agent_id: str) -> None:
    """Limpa apenas as abas de saída no início de uma rodada."""
    path = workbook_path(agent_id)
    book = load_workbook(path)
    _reset_output_sheets(book)
    _save_atomic(book, path)
    _refresh_csv_if_present(book, agent_id)


def append_result(
    agent_id: str,
    case: dict,
    *,
    status: str,
    answer: str = "",
    scores: dict | None = None,
    elapsed: float | None = None,
    reason: str = "",
    suggestion: str = "",
    error: str = "",
) -> None:
    path = workbook_path(agent_id)
    book = load_workbook(path)
    sheet = book["Métricas"]
    scores = scores or {}
    values = [scores.get(metric) for metric in METRICS]
    sheet.append(
        (
            case["id"],
            datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S"),
            status,
            case["question"],
            case["reference"],
            answer,
            case["mode"],
            *values,
            None,
            round(elapsed, 3) if elapsed is not None else None,
            reason,
            suggestion,
            error,
        )
    )
    row = sheet.max_row
    sheet.cell(row, 13, f'=IF(COUNT(H{row}:L{row})=0,"",AVERAGE(H{row}:L{row}))')
    for cell in sheet[row]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    _save_atomic(book, path)
    _refresh_csv_if_present(book, agent_id)
