"""A planilha é a entrada e recebe apenas os resultados da rodada atual."""

import csv
from types import SimpleNamespace

from openpyxl import load_workbook

from evals import support, workbook


def test_judge_accepts_alternative_completeness_key(monkeypatch):
    class FakeModel:
        def __init__(self, **_kwargs):
            pass

        def invoke(self, _messages):
            return SimpleNamespace(
                content='{"passed": true, "reason": "Equivalente", "suggestion": "", "scores": '
                '{"correcao_factual": 5, "completeude": 4, "relevancia": 5, '
                '"aderencia_instrucoes": 5, "clareza": 5}}'
            )

    monkeypatch.setattr(support, "ChatOpenAI", FakeModel)
    monkeypatch.setattr(
        support,
        "dotenv_values",
        lambda *_: {
            "OPENAI_MODEL": "test",
            "OPENAI_BASE_URL": "http://example.test",
            "OPENAI_API_KEY": "test",
            "OPENAI_TIMEOUT_SECONDS": "1",
        },
    )
    verdict = support.judge_answer("rsa-agent", "pergunta", "referência", "resposta")
    assert verdict["scores"]["completude"] == 4


def test_cases_survive_result_reset_and_failed_cases_are_recorded(tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "CASES_DIR", tmp_path)
    support._STARTED_AGENTS.clear()
    agent_id = "test-agent"
    original = {"id": "um", "question": "Pergunta", "reference": "Resposta", "mode": "semantic"}
    path = workbook.create_workbook(agent_id, [original])
    csv_path = path.with_suffix(".csv")
    csv_path.touch()

    book = load_workbook(path)
    book["Casos"].append(("dois", "Outra pergunta", "Outra resposta", "semantic"))
    book.save(path)
    cases = workbook.read_cases(agent_id)
    assert [case["id"] for case in cases] == ["um", "dois"]

    monkeypatch.setattr(support, "ask_agent", lambda *_: "Resposta incompleta")
    monkeypatch.setattr(
        support,
        "judge_answer",
        lambda *_: {
            "passed": False,
            "reason": "Faltou um fato.",
            "suggestion": "Inclua o fato esperado.",
            "scores": {metric: 2 for metric in workbook.METRICS},
        },
    )
    try:
        support.evaluate_case(agent_id, cases[0])
    except AssertionError:
        pass
    else:
        raise AssertionError("O caso reprovado deveria falhar no pytest")

    book = load_workbook(path)
    assert book["Métricas"]["C2"].value == "REPROVADO"
    assert book["Métricas"]["H2"].value == 2
    assert book["Métricas"]["P2"].value == "Inclua o fato esperado."
    assert book["Casos"].max_row == 3
    assert book["Resumo"]["B6"].value == "=IFERROR(B3/B2,0)"
    book.close()
    with csv_path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    assert rows[0]["status"] == "REPROVADO"

    workbook.reset_results(agent_id)
    book = load_workbook(path)
    assert book["Métricas"].max_row == 1
    assert book["Casos"].max_row == 3
    book.close()
    with csv_path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    assert rows[0]["status"] == ""
